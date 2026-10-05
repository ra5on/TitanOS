"""Validate real blank Office documents and the confined creation boundary."""
from io import BytesIO
import importlib.util
import json
import os
from pathlib import Path
import posixpath
import stat
import tempfile
import unittest
from unittest.mock import patch
from xml.etree import ElementTree as ET
from zipfile import ZIP_STORED, ZipFile

from titan.core import Error
from titan.document_templates import blank_document
from titan.files import operate
from titan.system_files import operate_system
from tests import test_admin_files as file_fixtures
from tests.test_lifecycle_http import HTTPFixture

REL = 'http://schemas.openxmlformats.org/package/2006/relationships'
OFFICE_REL = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'
CONTENT = 'http://schemas.openxmlformats.org/package/2006/content-types'
WORD = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
SHEET = 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'
ODF = 'urn:oasis:names:tc:opendocument:xmlns:office:1.0'
MANIFEST = 'urn:oasis:names:tc:opendocument:xmlns:manifest:1.0'


class DocumentFormatTests(unittest.TestCase):
    def test_every_office_template_is_a_valid_bounded_zip_without_macros_or_external_relationships(self):
        for kind in ('docx', 'xlsx', 'odt', 'ods'):
            with self.subTest(kind=kind), ZipFile(BytesIO(blank_document(kind, 'New.' + kind))) as archive:
                self.assertIsNone(archive.testzip())
                self.assertEqual(len(archive.namelist()), len(set(archive.namelist())))
                self.assertLess(sum(item.file_size for item in archive.infolist()), 32 * 1024)
                for name in archive.namelist():
                    self.assertFalse(name.startswith('/') or '..' in Path(name).parts)
                    self.assertNotIn('vba', name.lower())
                    self.assertNotIn('embeddings', name.lower())
                    if name.endswith(('.xml', '.rels')):
                        body = archive.read(name)
                        self.assertNotIn(b'<!DOCTYPE', body)
                        ET.fromstring(body)
                    if name.endswith('.rels'):
                        tree = ET.fromstring(archive.read(name))
                        ids = [item.attrib['Id'] for item in tree]
                        self.assertEqual(len(ids), len(set(ids)))
                        for relation in tree:
                            self.assertEqual(relation.tag, '{' + REL + '}Relationship')
                            self.assertNotEqual(relation.get('TargetMode'), 'External')
                            source = '' if name == '_rels/.rels' else name.replace('/_rels/', '/').removesuffix('.rels')
                            target = posixpath.normpath(posixpath.join(posixpath.dirname(source), relation.attrib['Target']))
                            self.assertIn(target, archive.namelist())

    def test_word_document_has_correct_content_type_namespace_body_and_page_section(self):
        with ZipFile(BytesIO(blank_document('docx', 'Brief.docx'))) as archive:
            types = ET.fromstring(archive.read('[Content_Types].xml'))
            entry = types.find('{%s}Override' % CONTENT)
            self.assertEqual(entry.attrib['PartName'], '/word/document.xml')
            self.assertTrue(entry.attrib['ContentType'].endswith('wordprocessingml.document.main+xml'))
            tree = ET.fromstring(archive.read('word/document.xml'))
            self.assertEqual(tree.tag, '{%s}document' % WORD)
            body = tree.find('{%s}body' % WORD)
            self.assertIsNotNone(body.find('{%s}p' % WORD))
            self.assertIsNotNone(body.find('{%s}sectPr/{%s}pgSz' % (WORD, WORD)))

    def test_excel_workbook_sheet_and_styles_are_connected_with_correct_namespaces(self):
        with ZipFile(BytesIO(blank_document('xlsx', 'Budget.xlsx'))) as archive:
            workbook = ET.fromstring(archive.read('xl/workbook.xml'))
            self.assertEqual(workbook.tag, '{%s}workbook' % SHEET)
            sheet = workbook.find('{%s}sheets/{%s}sheet' % (SHEET, SHEET))
            self.assertEqual(sheet.attrib['name'], 'Tabelle1')
            relation = sheet.attrib['{%s}id' % OFFICE_REL]
            relationships = ET.fromstring(archive.read('xl/_rels/workbook.xml.rels'))
            matching = next(item for item in relationships if item.attrib['Id'] == relation)
            self.assertEqual(matching.attrib['Target'], 'worksheets/sheet1.xml')
            self.assertTrue(matching.attrib['Type'].endswith('/worksheet'))
            worksheet = ET.fromstring(archive.read('xl/worksheets/sheet1.xml'))
            self.assertEqual(worksheet.tag, '{%s}worksheet' % SHEET)
            self.assertIsNotNone(worksheet.find('{%s}sheetData' % SHEET))
            styles = ET.fromstring(archive.read('xl/styles.xml'))
            self.assertEqual(styles.tag, '{%s}styleSheet' % SHEET)
            self.assertEqual(styles.find('{%s}cellXfs' % SHEET).attrib['count'], '1')

    def test_odf_mimetype_is_first_uncompressed_and_manifest_matches_content(self):
        for kind, suffix, body in (('odt', 'text', 'text'), ('ods', 'spreadsheet', 'spreadsheet')):
            with self.subTest(kind=kind), ZipFile(BytesIO(blank_document(kind, 'New.' + kind))) as archive:
                first = archive.infolist()[0]
                self.assertEqual(first.filename, 'mimetype')
                self.assertEqual(first.compress_type, ZIP_STORED)
                self.assertEqual(first.header_offset, 0)
                self.assertEqual(first.extra, b'')
                self.assertEqual(first.compress_size, first.file_size)
                mime = 'application/vnd.oasis.opendocument.' + suffix
                self.assertEqual(archive.read('mimetype').decode(), mime)
                manifest = ET.fromstring(archive.read('META-INF/manifest.xml'))
                entries = {item.attrib['{%s}full-path' % MANIFEST]: item.attrib['{%s}media-type' % MANIFEST] for item in manifest}
                self.assertEqual(entries['/'], mime)
                self.assertEqual(entries['content.xml'], 'text/xml')
                content = ET.fromstring(archive.read('content.xml'))
                self.assertIsNotNone(content.find('{%s}body/{%s}%s' % (ODF, ODF, body)))

    def test_unsupported_type_and_extension_mismatch_fail_before_any_file_write(self):
        for kind, name in (('doc', 'New.doc'), ('xls', 'New.xls'), ('docx', 'New.xlsx'), ('docx', 'New'), (None, 'New.docx'), ('xlsx', None)):
            with self.subTest(kind=kind, name=name), self.assertRaises(Error): blank_document(kind, name)
        with ZipFile(BytesIO(blank_document('docx', 'Brief.DOCX'))) as archive:
            self.assertIn('word/document.xml', archive.namelist())

    def test_office_templates_open_with_optional_real_python_readers(self):
        checked = []
        if importlib.util.find_spec('docx'):
            from docx import Document
            document = Document(BytesIO(blank_document('docx', 'New.docx')))
            self.assertEqual(len(document.paragraphs), 1)
            self.assertEqual(document.paragraphs[0].text, '')
            self.assertEqual(len(document.sections), 1)
            checked.append('docx')
        if importlib.util.find_spec('openpyxl'):
            from openpyxl import load_workbook
            workbook = load_workbook(BytesIO(blank_document('xlsx', 'New.xlsx')))
            self.assertEqual(workbook.sheetnames, ['Tabelle1'])
            self.assertIsNone(workbook.active['A1'].value)
            workbook.active['A1'] = 'Roundtrip'
            saved = BytesIO()
            workbook.save(saved)
            saved.seek(0)
            self.assertEqual(load_workbook(saved).active['A1'].value, 'Roundtrip')
            checked.append('xlsx')
        # The normal CI needs no optional document libraries. XML/OPC/ODF
        # contracts above always run; the bundled local runtime adds readers.
        self.assertIsInstance(checked, list)


class DocumentCreationTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory()
        self.addCleanup(self.temporary.cleanup)
        self.root = Path(self.temporary.name) / 'share'
        self.root.mkdir()

    def create(self, path, kind='docx'):
        return operate(self.root, 'create_document', path, document_type=kind)

    def test_actual_created_files_contain_valid_templates_owned_by_worker_uid(self):
        for kind in ('docx', 'xlsx', 'odt', 'ods'):
            name = 'Grüße.' + kind
            self.assertEqual(self.create(name, kind), {'ok': True})
            path = self.root / name
            self.assertEqual(path.stat().st_uid, os.geteuid())
            with ZipFile(path) as archive: self.assertIsNone(archive.testzip())

    def test_extension_failure_does_not_create_a_partial_file(self):
        with self.assertRaises(Error): self.create('Wrong.xlsx', 'docx')
        self.assertFalse((self.root / 'Wrong.xlsx').exists())

    def test_existing_file_folder_and_leaf_symlinks_are_never_overwritten(self):
        (self.root / 'old.docx').write_bytes(b'preserved')
        (self.root / 'folder.docx').mkdir()
        (self.root / 'link.docx').symlink_to('old.docx')
        (self.root / 'dangling.docx').symlink_to('missing.docx')
        for name in ('old.docx', 'folder.docx', 'link.docx', 'dangling.docx'):
            with self.subTest(name=name), self.assertRaises(Error): self.create(name)
        self.assertEqual((self.root / 'old.docx').read_bytes(), b'preserved')
        self.assertFalse((self.root / 'missing.docx').exists())
        self.assertTrue((self.root / 'link.docx').is_symlink())

    def test_creation_open_is_exclusive_nofollow_and_fsyncs_file_then_directory(self):
        synced = []
        original_sync = os.fsync
        def syncing(fd):
            synced.append(stat.S_IFMT(os.fstat(fd).st_mode))
            return original_sync(fd)
        with patch('titan.files.os.open', wraps=os.open) as opened, patch('titan.files.os.fsync', side_effect=syncing):
            self.create('New.docx')
        final = next(call for call in opened.call_args_list if call.args[0] == 'New.docx')
        self.assertEqual(final.args[1] & (os.O_EXCL | os.O_NOFOLLOW | os.O_CREAT), os.O_EXCL | os.O_NOFOLLOW | os.O_CREAT)
        self.assertIs(type(final.kwargs['dir_fd']), int)
        self.assertEqual(synced, [stat.S_IFREG, stat.S_IFDIR])

    def test_symlink_race_cannot_replace_an_outside_document(self):
        outside = Path(self.temporary.name) / 'outside.docx'
        outside.write_bytes(b'preserved')
        original_open = os.open
        def race(path, flags, *args, **kwargs):
            if path == 'New.docx' and flags & os.O_CREAT:
                (self.root / 'New.docx').symlink_to(outside)
            return original_open(path, flags, *args, **kwargs)
        with patch('titan.files.os.open', side_effect=race), self.assertRaises(Error) as caught:
            self.create('New.docx')
        self.assertEqual(caught.exception.status, 409)
        self.assertEqual(outside.read_bytes(), b'preserved')

    def test_renamed_parent_keeps_new_document_on_pinned_directory(self):
        (self.root / 'folder').mkdir()
        retained = self.root / 'retained'
        original_open = os.open
        def race(path, flags, *args, **kwargs):
            fd = original_open(path, flags, *args, **kwargs)
            if path == 'folder':
                (self.root / 'folder').rename(retained)
                (self.root / 'folder').mkdir()
            return fd
        with patch('titan.files.os.open', side_effect=race): self.create('folder/New.docx')
        self.assertTrue((retained / 'New.docx').is_file())
        self.assertFalse((self.root / 'folder/New.docx').exists())

    def test_failed_creation_removes_own_partial_inode_but_preserves_replacement(self):
        with patch('titan.files.os.fsync', side_effect=OSError('disk full')), self.assertRaises(OSError):
            self.create('New.docx')
        self.assertFalse((self.root / 'New.docx').exists())
        def replace(fd):
            (self.root / 'New.docx').unlink()
            (self.root / 'New.docx').write_bytes(b'other writer')
            raise OSError('disk full')
        with patch('titan.files.os.fsync', side_effect=replace), self.assertRaises(OSError): self.create('New.docx')
        self.assertEqual((self.root / 'New.docx').read_bytes(), b'other writer')

    def test_traversal_or_symlink_parent_cannot_create_outside_share(self):
        outside = Path(self.temporary.name) / 'outside'
        outside.mkdir()
        (self.root / 'escape').symlink_to(outside, target_is_directory=True)
        for name in ('../outside/New.docx', str(outside / 'New.docx'), 'escape/New.docx', 'bad\x00.docx'):
            with self.subTest(name=name), self.assertRaises((Error, OSError)): self.create(name)
        self.assertEqual(list(outside.iterdir()), [])

    def test_system_virtual_and_readonly_targets_refuse_before_creation(self):
        for directory in ('proc', 'sys', 'dev', 'usr', 'etc'):
            (self.root / directory).mkdir()
        for directory in ('proc', 'sys', 'dev'):
            with self.assertRaises(Error) as caught:
                operate_system(str(self.root), 'create_document', directory + '/New.docx', system_path_root=str(self.root), document_type='docx')
            self.assertEqual(caught.exception.status, 403)
            self.assertFalse((self.root / directory / 'New.docx').exists())
        with patch('titan.system_files.writable_path', return_value=False), self.assertRaises(Error) as caught:
            operate_system(str(self.root), 'create_document', 'usr/New.docx', system_path_root=str(self.root), document_type='docx')
        self.assertEqual(caught.exception.status, 403)
        self.assertFalse((self.root / 'usr/New.docx').exists())
        operate_system(str(self.root), 'create_document', 'etc/New.docx', system_path_root=str(self.root), document_type='docx')
        self.assertTrue((self.root / 'etc/New.docx').is_file())


class ShareDocumentHostTests(unittest.TestCase):
    setUp = file_fixtures.AdminFileHostTests.setUp
    tearDown = file_fixtures.AdminFileHostTests.tearDown
    execute_worker = file_fixtures.AdminFileHostTests.execute_worker

    def test_writer_creates_office_document_with_own_uid_and_groups(self):
        self.host.op_file('owner', 'private', 'create_document', 'New.xlsx', document_type='xlsx')
        request, values = self.calls[0]
        self.assertEqual(values['user'], 1234)
        self.assertEqual(values['group'], 4567)
        self.assertEqual(values['extra_groups'], [4567, 5678])
        with ZipFile(self.source / 'New.xlsx') as archive: self.assertIn('xl/workbook.xml', archive.namelist())

    def test_readonly_share_blocks_document_creation_before_worker(self):
        with self.assertRaises(Error) as caught:
            self.host.op_file('reader', 'private', 'create_document', 'New.docx', document_type='docx')
        self.assertEqual(caught.exception.status, 403)
        self.worker.assert_not_called()
        self.assertFalse((self.source / 'New.docx').exists())

    def test_admin_uses_verified_root_descriptor_and_rejects_extra_worker_arguments(self):
        self.host.op_admin_file('private', 'create_document', 'New.docx', document_type='docx')
        request, values = self.calls[0]
        self.assertIs(type(request['root']), int)
        self.assertEqual(values['pass_fds'], (request['root'],))
        self.assertEqual(values['user'], 0)
        self.worker.reset_mock()
        for extra in ({'root': '/'}, {'data': 'arbitrary'}, {'user': 'root'}):
            with self.assertRaises(Error):
                self.host.op_admin_file('private', 'create_document', 'Other.docx', document_type='docx', **extra)
        self.worker.assert_not_called()


class DocumentCreationHTTPTests(HTTPFixture, unittest.TestCase):
    def test_normal_user_dispatches_document_creation_under_own_account(self):
        body = {'share': 'public', 'path': 'New.docx', 'action': 'create_document', 'document_type': 'docx'}
        self.assertEqual(self.request('/api/files', body, actor='reader')[0], 200)
        self.agent.call.assert_called_once_with('file', user='reader', **body)

    def test_system_document_creation_is_admin_only_and_csrf_protected(self):
        body = {'share': '@system', 'path': 'etc/New.docx', 'action': 'create_document', 'document_type': 'docx'}
        self.assertEqual(self.request('/api/files', body, actor='reader')[0], 403)
        self.assertEqual(self.request('/api/files', body, csrf='wrong')[0], 403)
        self.assertEqual(self.request('/api/files', body, headers={'Origin': 'https://foreign.example'})[0], 403)
        self.agent.call.assert_not_called()
        self.assertEqual(self.request('/api/files', body)[0], 200)
        self.agent.call.assert_called_once_with('system_file', action='create_document', path='etc/New.docx', document_type='docx')

    def test_privileged_worker_flags_cannot_be_injected_over_http(self):
        body = {'share': 'public', 'path': 'New.docx', 'action': 'create_document', 'document_type': 'docx'}
        for extra in ({'root': '/'}, {'user': 'root'}, {'system': True}, {'canonicalized': True}):
            self.assertEqual(self.request('/api/files', {**body, **extra}, actor='reader')[0], 400)
        self.agent.call.assert_not_called()


if __name__ == '__main__': unittest.main()
